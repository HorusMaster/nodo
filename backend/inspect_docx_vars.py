from docx import Document
import re
import os

template_path = "pantillas/LW/LW CONTRATO.docx"

if not os.path.exists(template_path):
    print(f"Error: {template_path} not found")
    exit(1)

doc = Document(template_path)

print(f"Variables found in {template_path}:")
print("-" * 30)

variables = set()
pattern = re.compile(r"\{\{([^}]+)\}\}")

def extract_vars(text):
    matches = pattern.findall(text)
    for match in matches:
        variables.add(match.strip())

for paragraph in doc.paragraphs:
    extract_vars(paragraph.text)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                extract_vars(paragraph.text)

for var in sorted(variables):
    print(f"Found variable: {var}")

print("-" * 30)
