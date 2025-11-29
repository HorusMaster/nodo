import sys
import os
from docx import Document
import io

# Add backend to path
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from pantillas.SUSHOP.contrato_generator import generate_contract_docx

def test_integration():
    # Create a dummy template
    template_path = "dummy_template.docx"
    doc = Document()
    doc.add_paragraph("Contract for {{receptor}}")
    doc.add_paragraph("Missing: {{missing_var}}")
    doc.save(template_path)
    
    # Test data
    data = {"receptor": "Jane Doe"}
    
    try:
        # Run generator
        pdf_stream = generate_contract_docx(data, template_path)
        
        # Read result
        result_doc = Document(pdf_stream)
        text = "\n".join([p.text for p in result_doc.paragraphs])
        
        print("Result Text:")
        print(text)
        
        if "Contract for Jane Doe" in text and "Missing: XXXXXXXXXX" in text:
            print("Integration Test PASSED")
        else:
            print("Integration Test FAILED")
            
    finally:
        if os.path.exists(template_path):
            os.remove(template_path)

if __name__ == "__main__":
    test_integration()
