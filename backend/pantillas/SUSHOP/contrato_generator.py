"""
Generador de contratos DOCX para LW usando python-docx.
Rellena la plantilla LW CONTRATO.docx con datos de n8n.
Los campos faltantes se rellenan con "XXXXXXXXXX".

NOTA: Esta versión usa python-docx directamente en lugar de docxtpl
debido a un bug en docxtpl 0.20.2 con el manejo de rutas.
"""

import io
import os
from pathlib import Path
from docx import Document
from datetime import datetime


import re

def replace_text_in_paragraph(paragraph, replacements):
    """
    Reemplaza texto en un párrafo manteniendo el formato lo mejor posible.
    Maneja casos donde el placeholder está dividido en múltiples 'runs'.
    """
    # 1. Intento rápido: Reemplazo directo en runs (conserva formato perfecto)
    # Expandimos las claves para incluir versiones con espacios: {{ key }} y {{key}}
    extended_replacements = {}
    for key, value in replacements.items():
        extended_replacements[key] = value
        # Si la clave es {{key}}, agregar {{ key }}
        if key.startswith('{{') and key.endswith('}}'):
            inner = key[2:-2]
            extended_replacements[f'{{{{ {inner} }}}}'] = value
            extended_replacements[f'{{{{ {inner}}}}}'] = value
            extended_replacements[f'{{{{{inner} }}}}'] = value

    # Bandera para saber si hicimos algún cambio
    replaced = False
    
    for key, value in extended_replacements.items():
        if key in paragraph.text:
            # Intentar reemplazar en cada run individualmente
            for run in paragraph.runs:
                if key in run.text:
                    run.text = run.text.replace(key, str(value))
                    run.bold = True  # Aplicar negrita
                    replaced = True
    
    # 2. Si el texto está en el párrafo pero no se reemplazó (split runs),
    # usamos una estrategia más agresiva: reemplazo en el texto del párrafo.
    # Esto puede perder formato mixto (negrita/normal en la misma línea),
    # pero garantiza que el dato se rellene.
    if not replaced:
        original_text = paragraph.text
        new_text = original_text
        
        # Usar regex para encontrar patrones {{ key }} con cualquier cantidad de espacios
        # Iteramos sobre las claves originales (sin espacios extra)
        for key, value in replacements.items():
            if key.startswith('{{') and key.endswith('}}'):
                inner = re.escape(key[2:-2])
                pattern = rf"{{{{\s*{inner}\s*}}}}"
                if re.search(pattern, new_text):
                    new_text = re.sub(pattern, str(value), new_text)
                    replaced = True
        
        # Si hubo cambios, actualizamos el párrafo
        if replaced and new_text != original_text:
            # Limpiar todos los runs y poner el nuevo texto en el primero (o uno nuevo)
            # Para intentar conservar el formato del primer run:
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                paragraph.runs[0].bold = True  # Aplicar negrita al run modificado
                # Eliminar el texto de los demás runs para evitar duplicados
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                run = paragraph.add_run(new_text)
                run.bold = True  # Aplicar negrita al nuevo run


def replace_text_in_table(table, replacements):
    """Reemplaza texto en una tabla."""
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_text_in_paragraph(paragraph, replacements)


def generate_contract_docx(data_item, template_path):
    """
    Genera un contrato DOCX rellenado con los datos proporcionados.
    Escanea la plantilla en busca de variables {{variable}} y las rellena
    con los datos de n8n o 'XXXXXXXXXX' si no existen.
    
    Args:
        data_item (dict): Datos del contrato desde n8n
        template_path (str): Ruta a la plantilla DOCX
        
    Returns:
        io.BytesIO: Stream de bytes del documento DOCX generado
    """
    
    # Convertir a Path object
    template_path = Path(template_path)
    if not template_path.is_absolute():
        template_path = template_path.resolve()
    
    # Cargar el documento
    doc = Document(str(template_path))
    
    # 1. Pre-procesar datos (Contexto)
    context = data_item.copy()
    
    # Manejar fecha y hora derivadas
    fecha_str = context.get('fecha')
    if not fecha_str:
        fecha_str = context.get('fecha_hora_emision')
    if not fecha_str:
        fecha_str = context.get('fecha/hora_emision')
    
    if fecha_str:
        try:
            # Normalizar separador T a espacio
            fecha_val = str(fecha_str).replace('T', ' ')
            context['fecha/hora_emision'] = fecha_val
            
            if ' ' in fecha_val:
                parts = fecha_val.split(' ')
                if 'fecha' not in context: context['fecha'] = parts[0]
                if 'hora_emision' not in context: context['hora_emision'] = parts[1]
            else:
                if 'fecha' not in context: context['fecha'] = fecha_val
        except:
            pass

    # 2. Escanear el documento para encontrar todas las variables {{key}}
    found_keys = set()
    
    def find_keys_in_text(text):
        return set(re.findall(r'\{\{\s*(\w+)\s*\}\}', text))

    for paragraph in doc.paragraphs:
        found_keys.update(find_keys_in_text(paragraph.text))
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    found_keys.update(find_keys_in_text(paragraph.text))

    # 3. Construir diccionario de reemplazos
    replacements = {}
    for key in found_keys:
        # Obtener valor del contexto o usar default
        val = context.get(key, 'XXXXXXXXXX')
        # La función replace_text_in_paragraph espera claves con llaves {{key}}
        replacements[f'{{{{{key}}}}}'] = val
    
    # 4. Ejecutar reemplazos
    # Reemplazar en todos los párrafos
    for paragraph in doc.paragraphs:
        replace_text_in_paragraph(paragraph, replacements)
    
    # Reemplazar en todas las tablas
    for table in doc.tables:
        replace_text_in_table(table, replacements)
    
    # Guardar en un BytesIO stream
    output_stream = io.BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)
    
    return output_stream


def generate_contract_bytes(data_item, template_path):
    """
    Alias para mantener consistencia con otros generadores.
    Genera un contrato DOCX y retorna los bytes.
    
    Args:
        data_item (dict): Datos del contrato desde n8n
        template_path (str): Ruta a la plantilla DOCX
        
    Returns:
        io.BytesIO: Stream de bytes del documento DOCX generado
    """
    return generate_contract_docx(data_item, template_path)
