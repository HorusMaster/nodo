from docx import Document
import re
import io

def replace_placeholders(doc, data):
    """
    Replaces {{key}} placeholders in the document with values from data.
    Defaults to 'XXXXXXXXXX' if key is missing in data.
    """
    def replace_in_text(text):
        # Find all {{key}} patterns
        matches = re.findall(r'\{\{\s*(\w+)\s*\}\}', text)
        for key in matches:
            val = data.get(key, 'XXXXXXXXXX')
            # Replace {{ key }} or {{key}} with value
            # We use a regex to match the specific occurrence including braces and spaces
            pattern = r'\{\{\s*' + re.escape(key) + r'\s*\}\}'
            text = re.sub(pattern, str(val), text)
        return text

    for paragraph in doc.paragraphs:
        if '{{' in paragraph.text:
            # Simple replacement strategy: replace in the full text and reset runs
            # This might lose formatting within the paragraph but ensures replacement works
            new_text = replace_in_text(paragraph.text)
            if new_text != paragraph.text:
                paragraph.text = new_text

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if '{{' in paragraph.text:
                        new_text = replace_in_text(paragraph.text)
                        if new_text != paragraph.text:
                            paragraph.text = new_text

# Create a dummy document
doc = Document()
doc.add_paragraph("Contract for {{receptor}}")
doc.add_paragraph("Date: {{fecha}}")
doc.add_paragraph("Unknown: {{missing_var}}")

# Test data
data = {
    "receptor": "John Doe",
    "fecha": "2023-11-29"
}

# Run replacement
replace_placeholders(doc, data)

# Verify results
print("Paragraphs:")
for p in doc.paragraphs:
    print(p.text)

# Expected:
# Contract for John Doe
# Date: 2023-11-29
# Unknown: XXXXXXXXXX
